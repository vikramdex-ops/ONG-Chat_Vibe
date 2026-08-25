from pathlib import Path
from typing import List, Dict, Any
from app.core.config import IMAGES_DIR
from app.core.logging_service import app_logger
from app.services.vector_db import sanitize_filename

try:
    import docx
except ImportError:
    docx = None


def parse_docx_document(file_path: Path) -> List[Dict[str, Any]]:
    """
    Extracts paragraphs, table cells, and inline images from a DOCX file.
    Represents the document as a single conceptual page (page=1) to maintain SQA-O&G compatibility.
    Returns list of dicts: [{'page': 1, 'text': str, 'image_paths': List[str]}]
    """
    if not docx:
        app_logger.error("DOCXParser", "python-docx is not installed. Cannot parse DOCX.", document=file_path.name)
        raise RuntimeError("python-docx library is required to process DOCX files.")

    base_name = sanitize_filename(file_path.stem)
    all_doc_text = []
    image_paths = []
    doc_img_index = 0

    try:
        document = docx.Document(file_path)
        app_logger.debug("DOCXParser", f"Parsing DOCX '{file_path.name}'", document=file_path.name)

        # 1. Paragraphs
        for para in document.paragraphs:
            if para.text.strip():
                all_doc_text.append(para.text.strip())

        # 2. Tables
        for table in document.tables:
            for row in table.rows:
                row_texts = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        row_texts.append(cell_text)
                if row_texts:
                    all_doc_text.append(" | ".join(row_texts))

        # 3. Inline Images
        for shape in document.inline_shapes:
            try:
                if hasattr(shape, "_inline") and hasattr(shape._inline, "graphic"):
                    blip = shape._inline.graphic.graphicData.pic.blipFill.blip
                    rId = blip.embed
                    image_part = document.part.related_parts[rId]
                    image_bytes = image_part.blob
                    image_ext = image_part.ext or "png"

                    img_filename = f"{base_name}_img_{doc_img_index}.{image_ext}"
                    save_path = IMAGES_DIR / img_filename
                    with open(save_path, "wb") as img_file:
                        img_file.write(image_bytes)
                    image_paths.append(str(save_path))
                    doc_img_index += 1
            except Exception as img_ex:
                app_logger.debug("DOCXParser", f"Image extraction note: {img_ex}", document=file_path.name)

        full_doc_text = "\n\n".join(all_doc_text).strip()
        if full_doc_text or image_paths:
            return [{
                "page": 1,
                "text": full_doc_text,
                "image_paths": image_paths
            }]
        return []
    except Exception as e:
        app_logger.error("DOCXParser", f"Failed to parse DOCX {file_path.name}: {e}", document=file_path.name)
        raise
