from pathlib import Path
from typing import List, Dict, Any
from app.core.config import IMAGES_DIR
from app.core.logging_service import app_logger
from app.services.vector_db import sanitize_filename

try:
    import pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    pptx = None


def parse_pptx_document(file_path: Path) -> List[Dict[str, Any]]:
    """
    Extracts text and images from a PowerPoint (.pptx) file slide-by-slide.
    Returns list of dicts: [{'page': int (slide_num), 'text': str, 'image_paths': List[str]}, ...]
    """
    if not pptx:
        app_logger.error("PPTXParser", "python-pptx is not installed. Cannot parse PPTX.", document=file_path.name)
        raise RuntimeError("python-pptx library is required to process PowerPoint files.")

    base_name = sanitize_filename(file_path.stem)
    page_data_list = []

    try:
        presentation = pptx.Presentation(file_path)
        app_logger.debug("PPTXParser", f"Parsing PPTX '{file_path.name}' ({len(presentation.slides)} slides)", document=file_path.name)

        for i, slide in enumerate(presentation.slides):
            slide_num = i + 1
            slide_text = []
            image_paths = []
            shape_img_index = 0

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.text.strip():
                                slide_text.append(run.text.strip())

                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image = shape.image
                        image_bytes = image.blob
                        image_ext = image.ext or "png"
                        img_filename = f"{base_name}_slide_{slide_num}_img_{shape_img_index}.{image_ext}"
                        save_path = IMAGES_DIR / img_filename
                        with open(save_path, "wb") as img_file:
                            img_file.write(image_bytes)
                        image_paths.append(str(save_path))
                        shape_img_index += 1
                    except Exception as img_ex:
                        app_logger.warning("PPTXParser", f"Slide {slide_num}: Error extracting image: {img_ex}", document=file_path.name)

            full_slide_text = " ".join(slide_text).strip()
            if full_slide_text or image_paths:
                page_data_list.append({
                    "page": slide_num,
                    "text": full_slide_text,
                    "image_paths": image_paths
                })

        return page_data_list
    except Exception as e:
        app_logger.error("PPTXParser", f"Failed to parse PPTX {file_path.name}: {e}", document=file_path.name)
        raise
